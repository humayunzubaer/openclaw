"""
Excel Engine — বুদ্ধিমান স্প্রেডশিট রিডার
==========================================

বাস্তব অডিট ফাইলগুলো নোংরা হয়:
- হেডার ৫ম বা ১০ম সারিতে থাকে
- Merged Cell থাকে
- Subtotal সারি মাঝখানে থাকে
- একাধিক Sheet, কিছু Hidden
- Formula থাকে (মান নয়)
- বাংলা ও ইংরেজি মিশ্রিত কলাম নাম

এই ইঞ্জিন সেসব স্বয়ংক্রিয়ভাবে সামলায়।
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

import numpy as np
import pandas as pd
from utils.logger import logger


# ==========================================================
# কলাম নামের অভিধান — বাস্তব ফাইলে যেসব নাম দেখা যায়
# ==========================================================

COLUMN_SYNONYMS: dict[str, list[str]] = {
    # --- সিরিয়াল ---
    "serial_no": [
        "sl", "sl.", "sl no", "sl. no", "slno", "serial", "serial no",
        "si", "si no", "s/n", "sn", "#", "ক্রমিক", "ক্রমিক নং", "ক্র নং",
    ],
    # --- HS Code ---
    "hs_code": [
        "hs code", "hscode", "h.s. code", "h s code", "hs", "hs-code",
        "harmonized code", "tariff code", "hs code no", "hs code number",
        "এইচএস কোড", "এইচ এস কোড", "শুল্ক কোড",
    ],
    # --- পণ্যের নাম ---
    "item_name": [
        "description", "item description", "description of goods",
        "goods description", "item name", "product name", "particulars",
        "name of item", "name of goods", "commodity", "item",
        "description of item", "material description", "raw material",
        "পণ্যের নাম", "পণ্যের বিবরণ", "মালামালের বিবরণ", "বিবরণ",
    ],
    "commercial_name": [
        "commercial description", "commercial name", "trade name",
        "brand name", "local name", "commercial invoice description",
        "বাণিজ্যিক নাম",
    ],
    # --- পরিমাণ ---
    "quantity": [
        "qty", "quantity", "qnty", "quantiy", "total qty", "net qty",
        "imported qty", "import quantity", "approved qty", "quantity imported",
        "পরিমাণ", "মোট পরিমাণ",
    ],
    "unit": [
        "unit", "uom", "u/m", "unit of measurement", "measurement unit",
        "unit name", "একক",
    ],
    "approved_quantity": [
        "approved quantity", "approved qty", "entitled qty", "entitlement qty",
        "permitted quantity", "allowed qty", "annual entitlement",
        "অনুমোদিত পরিমাণ",
    ],
    # ★ প্রাপ্যতা কলাম — শিরোনামে মেয়াদ উল্লেখ থাকে
    "entitled_quantity": [
        "প্রাপ্যতা", "আমদানি প্রাপ্যতা", "প্রদত্ত", "প্রস্তাবিত",
        "প্রদত্ত/প্রস্তাবিত", "প্রাপ্যতার পরিমাণ", "নির্ধারিত প্রাপ্যতা",
        "import entitlement", "entitlement", "entitled quantity",
    ],
    # ★ সমাপনী মজুত — পরবর্তী নিরীক্ষার প্রারম্ভিক জের
    "closing_stock": [
        "সমাপনী মজুত", "সমাপনী জের", "সমাপনী স্থিতি", "মজুত",
        "closing stock", "closing balance", "balance stock",
        "stock in hand", "warehouse stock", "মজুদ",
    ],
    "opening_stock": [
        "প্রারম্ভিক জের", "প্রারম্ভিক মজুত", "প্রারম্ভিক স্থিতি",
        "opening stock", "opening balance", "b/f", "brought forward",
    ],
    # ★ মেশিনের বার্ষিক উৎপাদন ক্ষমতা — বিধি ১১(১) যাচাইয়ের জন্য
    "capacity_100": [
        "উৎপাদন ক্ষমতার ১০০%", "উৎপাদন ক্ষমতা ১০০%", "মেশিনের উৎপাদন ক্ষমতা",
        "বার্ষিক উৎপাদন ক্ষমতা", "১০০%", "100%", "capacity 100",
        "production capacity 100", "annual production capacity",
    ],
    "capacity_80": [
        "উৎপাদন ক্ষমতার ৮০%", "উৎপাদন ক্ষমতা ৮০%", "ক্ষমতার ৮০%",
        "৮০%", "80%", "capacity 80", "production capacity 80",
    ],
    "capacity_60": [
        "উৎপাদন ক্ষমতার ৬০%", "উৎপাদন ক্ষমতা ৬০%", "ক্ষমতার ৬০%",
        "৬০%", "60%", "capacity 60", "production capacity 60",
    ],
    # ★ ওয়্যারহাউসের ধারণ ক্ষমতা — ২০২৬ সালের নূতন নিয়মে ব্যবহৃত
    "warehouse_capacity": [
        "ওয়্যারহাউসের ধারণ ক্ষমতা", "ধারণ ক্ষমতা", "গুদামের ধারণ ক্ষমতা",
        "warehouse holding capacity", "holding capacity", "storage capacity",
    ],
    # ★ এককালীন বন্ডিং ক্যাপাসিটি
    "bonding_capacity": [
        "বন্ডিং ক্যাপাসিটি", "এককালীন বন্ডিং ক্যাপাসিটি", "বন্ড ক্ষমতা",
        "ধারণক্ষমতা", "bonding capacity", "warehouse capacity",
        "one time bonding capacity", "bond capacity",
    ],
    # ★ স্থানীয় ক্রয় / ডিমড আমদানি
    "local_purchase": [
        "স্থানীয় ক্রয়", "স্থানীয় সংগ্রহ", "local purchase",
        "local procurement", "deemed import", "local supply",
        "back to back", "local sourcing",
    ],
    # --- মূল্য ---
    "unit_price": [
        "unit price", "rate", "price", "unit rate", "u/price",
        "price per unit", "unit value", "একক মূল্য",
    ],
    "value_usd": [
        "value", "total value", "amount", "cif value", "fob value",
        "value usd", "usd value", "value in usd", "total amount",
        "assessable value", "invoice value", "মূল্য", "মোট মূল্য",
    ],
    "value_bdt": [
        "value bdt", "bdt value", "value in bdt", "value in taka",
        "assessable value bdt", "taka value", "টাকা", "মূল্য (টাকা)",
    ],
    # --- বিল ---
    "bill_number": [
        "be no", "b/e no", "be number", "bill of entry", "bill of entry no",
        "bill no", "bill number", "im no", "im-4 no", "im4 no",
        "declaration no", "c no", "বিল নং", "বিই নং",
    ],
    "bill_date": [
        "be date", "b/e date", "bill date", "date", "bill of entry date",
        "assessment date", "im date", "তারিখ", "বিল তারিখ",
    ],
    "bill_type": [
        "type", "bill type", "be type", "im type", "declaration type",
    ],
    # --- শুল্ক ---
    "duty_rate": ["cd", "cd rate", "customs duty", "duty rate", "cd%", "শুল্ক হার"],
    "vat_rate": ["vat", "vat rate", "vat%", "ভ্যাট হার"],
    "at_rate": ["at", "at rate", "advance tax", "at%"],
    "rd_rate": ["rd", "rd rate", "regulatory duty", "rd%"],
    "sd_rate": ["sd", "sd rate", "supplementary duty", "sd%"],
    "ait_rate": ["ait", "ait rate", "advance income tax", "ait%"],
    "duty_paid": ["duty paid", "cd paid", "customs duty paid", "total duty",
                  "cd amount", "customs duty amount", "শুল্ক", "আরোপিত শুল্ক"],
    "vat_paid": ["vat paid", "vat amount", "total vat", "মূসক", "ভ্যাট"],
    "rd_paid": ["rd paid", "rd amount", "regulatory duty paid",
                "regulatory duty amount", "নিয়ন্ত্রণমূলক শুল্ক"],
    "sd_paid": ["sd paid", "sd amount", "supplementary duty paid",
                "supplementary duty amount", "সম্পূরক শুল্ক"],
    "at_paid": ["at paid", "at amount", "advance tax paid",
                "advance tax amount", "অগ্রিম কর"],
    "ait_paid": ["ait paid", "ait amount", "advance income tax paid",
                 "advance income tax amount", "অগ্রিম আয়কর"],
    # ★ MIS-এ উল্লিখিত মোট শুল্ক-কর — বন্ড সুবিধা বাতিলে সম্পূর্ণটাই দাবিযোগ্য
    "total_tax": [
        "total tax", "total duty tax", "total duty and tax",
        "total payable", "total revenue", "grand total tax",
        "মোট শুল্ক কর", "মোট শুল্ক-কর", "মোট রাজস্ব", "সর্বমোট শুল্ক",
        "duty tax total", "total assessed",
    ],
    # --- অন্যান্য ---
    "supplier": [
        "supplier", "supplier name", "exporter", "seller", "shipper",
        "consignor", "সরবরাহকারী",
    ],
    "country": [
        "country", "origin", "country of origin", "coo", "source country",
        "উৎস দেশ",
    ],
    "lc_number": ["lc no", "l/c no", "lc number", "letter of credit"],
    "bl_number": ["bl no", "b/l no", "bill of lading", "awb no"],
    "port": ["port", "port of entry", "customs station", "station"],
    "exchange_rate": ["exchange rate", "ex rate", "conversion rate", "rate of exchange"],
}

# উল্টো ম্যাপ তৈরি (দ্রুত lookup এর জন্য)
_SYNONYM_LOOKUP: dict[str, str] = {}
for canonical, variants in COLUMN_SYNONYMS.items():
    for v in variants:
        _SYNONYM_LOOKUP[v] = canonical


# বাংলা → ইংরেজি অঙ্ক রূপান্তর মানচিত্র
_BN_DIGIT_MAP = str.maketrans("০১২৩৪৫৬৭৮৯", "0123456789")


def _compact(s: str) -> str:
    """স্পেস, %, বিন্দু, হাইফেন সরিয়ে সংক্ষিপ্ত রূপ — 'CD %' → 'cd'"""
    return re.sub(r"[\s%./\-_]", "", s.lower())


# সংক্ষিপ্ত রূপের ম্যাপ — 'cd', 'at', 'rd', 'sd' এর মতো ছোট কোডের জন্য
_COMPACT_LOOKUP: dict[str, str] = {}
for _syn, _canon in _SYNONYM_LOOKUP.items():
    _c = _compact(_syn)
    if _c and _c not in _COMPACT_LOOKUP:
        _COMPACT_LOOKUP[_c] = _canon


# ==========================================================
# Subtotal / Total সারি শনাক্তকরণের শব্দ
# ==========================================================
TOTAL_KEYWORDS = {
    "total", "sub total", "subtotal", "sub-total", "grand total",
    "g. total", "g total", "sum", "মোট", "সর্বমোট", "উপমোট",
}


@dataclass
class SheetProfile:
    """একটি Sheet এর বিশ্লেষণ ফলাফল"""
    sheet_name: str
    is_hidden: bool = False
    header_row: Optional[int] = None
    data_start_row: Optional[int] = None
    data_end_row: Optional[int] = None
    total_rows: int = 0
    total_cols: int = 0
    detected_columns: dict[str, str] = field(default_factory=dict)  # canonical → actual
    unmapped_columns: list[str] = field(default_factory=list)
    subtotal_rows: list[int] = field(default_factory=list)
    merged_ranges: list[str] = field(default_factory=list)
    merged_parsed: list[dict] = field(default_factory=list)  # raw df স্থানাঙ্কে
    column_positions: dict[str, int] = field(default_factory=dict)
    has_formulas: bool = False
    confidence: float = 0.0
    likely_type: str = "unknown"  # entitlement | import | consumption | export | unknown


@dataclass
class ExcelReadResult:
    """সম্পূর্ণ ফাইল পড়ার ফলাফল"""
    file_path: str
    file_name: str
    sheets: list[SheetProfile] = field(default_factory=list)
    dataframes: dict[str, pd.DataFrame] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def sheet_names(self) -> list[str]:
        return [s.sheet_name for s in self.sheets]

    def get_sheet(self, name: str) -> Optional[pd.DataFrame]:
        return self.dataframes.get(name)

    def best_sheet(self, sheet_type: str = None) -> Optional[str]:
        """সবচেয়ে সম্ভাব্য Sheet খুঁজে দেয়"""
        candidates = [s for s in self.sheets if s.header_row is not None]
        if sheet_type:
            typed = [s for s in candidates if s.likely_type == sheet_type]
            if typed:
                candidates = typed
        if not candidates:
            return None
        return max(candidates, key=lambda s: (s.confidence, s.total_rows)).sheet_name


# ==========================================================
# Helper Functions
# ==========================================================

def normalize_text(text: Any) -> str:
    """টেক্সট স্বাভাবিকীকরণ — তুলনার জন্য"""
    if text is None or (isinstance(text, float) and np.isnan(text)):
        return ""
    s = str(text)
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("\n", " ").replace("\r", " ").replace("\t", " ")
    s = re.sub(r"[^\w\s\u0980-\u09FF%/.-]", " ", s)  # বাংলা ইউনিকোড রক্ষা
    s = re.sub(r"\s+", " ", s)
    return s.strip().lower()


def is_total_row(row: pd.Series) -> bool:
    """এই সারিটি কি Total/Subtotal?"""
    for val in row:
        n = normalize_text(val)
        if not n:
            continue
        # শুধু "total" শব্দটি থাকলেই যথেষ্ট নয়, cell টি ছোট হতে হবে
        if len(n) <= 25:
            for kw in TOTAL_KEYWORDS:
                if n == kw or n.startswith(kw + " ") or n.endswith(" " + kw):
                    return True
    return False


def clean_number(value: Any) -> Optional[float]:
    """
    যেকোনো ফরম্যাটের সংখ্যা পরিষ্কার করে float বানায়।
    সামলায়: "1,234.56", "১২৩৪", "(500)", "USD 1200", "12 000", "-"
    """
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and np.isnan(value):
            return None
        return float(value)

    s = str(value).strip()
    if not s or s in {"-", "--", "n/a", "N/A", "nil", "NIL", ""}:
        return None

    # বাংলা সংখ্যা → ইংরেজি
    bn_digits = "০১২৩৪৫৬৭৮৯"
    for i, d in enumerate(bn_digits):
        s = s.replace(d, str(i))

    # বন্ধনী মানে ঋণাত্মক (accounting format)
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative = True
        s = s[1:-1]

    # মুদ্রা প্রতীক ও অক্ষর সরাও
    s = re.sub(r"[^\d.,\-+eE]", "", s)
    if not s:
        return None

    # কমা হ্যান্ডলিং — "1,234.56" vs "1.234,56" (ইউরোপীয়)
    if "," in s and "." in s:
        if s.rindex(",") > s.rindex("."):
            s = s.replace(".", "").replace(",", ".")   # ইউরোপীয়
        else:
            s = s.replace(",", "")                      # স্ট্যান্ডার্ড
    elif "," in s:
        parts = s.split(",")
        # "1,234" → হাজার বিভাজক | "12,5" → দশমিক
        if len(parts[-1]) == 3 and len(parts) > 1:
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")

    try:
        result = float(s)
        return -result if negative else result
    except (ValueError, TypeError):
        return None


def clean_hs_code(value: Any) -> Optional[str]:
    """
    HS কোড স্বাভাবিকীকরণ।
    "5208.11.00" → "5208.11.00"
    "52081100"   → "5208.11.00"
    "5208 11 00" → "5208.11.00"
    Excel এ 5208110000 float হিসেবে আসতে পারে
    """
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.lower() in {"nan", "none", "-"}:
        return None

    # float হিসেবে এলে (যেমন 52081100.0)
    if s.endswith(".0"):
        s = s[:-2]

    # শুধু সংখ্যা রাখো
    digits = re.sub(r"\D", "", s)
    if len(digits) < 4:
        return None

    # ফরম্যাট: XXXX.XX.XX
    if len(digits) >= 8:
        return f"{digits[:4]}.{digits[4:6]}.{digits[6:8]}"
    elif len(digits) >= 6:
        return f"{digits[:4]}.{digits[4:6]}"
    else:
        return digits[:4]


def parse_merged_range(rng: str) -> Optional[dict]:
    """
    "F7:F10" → {"col_start":5,"col_end":5,"row_start":6,"row_end":9}
    (০-ভিত্তিক raw DataFrame স্থানাঙ্ক)
    """
    m = re.match(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$", str(rng).strip().upper())
    if not m:
        return None

    def col_to_idx(letters: str) -> int:
        n = 0
        for ch in letters:
            n = n * 26 + (ord(ch) - 64)
        return n - 1

    return {
        "col_start": col_to_idx(m.group(1)),
        "row_start": int(m.group(2)) - 1,
        "col_end": col_to_idx(m.group(3)),
        "row_end": int(m.group(4)) - 1,
    }


def match_column(header_text: str) -> Optional[str]:
    """একটি কলাম হেডারকে canonical নামে ম্যাপ করে"""
    n = normalize_text(header_text)
    if not n:
        return None

    # ০. উৎপাদন ক্ষমতার শতাংশ — সর্বোচ্চ অগ্রাধিকার
    #    ("মেশিনের উৎপাদন ক্ষমতার ৮০%" যেন capacity_100 এ না যায়)
    if any(k in n for k in ("ক্ষমতা", "capacity")):
        pct = re.search(r"(\d{2,3})\s*%", n.translate(_BN_DIGIT_MAP))
        if pct:
            val = pct.group(1)
            if val in {"100", "80", "60"}:
                return f"capacity_{val}"

    # ১. সরাসরি মিল
    if n in _SYNONYM_LOOKUP:
        return _SYNONYM_LOOKUP[n]

    # ২. সংক্ষিপ্ত রূপে মিল — "CD %" → "cd", "B/E No." → "beno"
    c = _compact(n)
    if c in _COMPACT_LOOKUP:
        return _COMPACT_LOOKUP[c]

    # ৩. আংশিক মিল — দীর্ঘতম synonym আগে (ভুল মিল এড়াতে)
    for syn in sorted(_SYNONYM_LOOKUP.keys(), key=len, reverse=True):
        if len(syn) < 3:
            continue
        if syn in n or n in syn:
            return _SYNONYM_LOOKUP[syn]

    return None


# ==========================================================
# মূল Excel Engine
# ==========================================================

class ExcelEngine:
    """
    বুদ্ধিমান Excel/CSV রিডার।

    ব্যবহার:
        engine = ExcelEngine()
        result = engine.read("entitlement.xlsx")
        df = result.get_sheet(result.best_sheet())
    """

    MAX_HEADER_SCAN_ROWS = 30   # প্রথম কত সারিতে হেডার খুঁজবে
    MIN_HEADER_MATCHES = 2      # কমপক্ষে কয়টি কলাম মিললে হেডার ধরবে

    def __init__(self, verbose: bool = True):
        self.verbose = verbose

    # ------------------------------------------------------
    def read(self, file_path: str | Path) -> ExcelReadResult:
        """ফাইল পড়ো এবং বিশ্লেষণ করো"""
        path = Path(file_path)
        result = ExcelReadResult(file_path=str(path), file_name=path.name)

        if not path.exists():
            result.errors.append(f"ফাইল পাওয়া যায়নি: {path}")
            return result

        ext = path.suffix.lower()

        try:
            if ext == ".csv":
                self._read_csv(path, result)
            elif ext in {".xlsx", ".xlsm", ".xls"}:
                self._read_excel(path, result)
            else:
                result.errors.append(f"অসমর্থিত ফাইল ফরম্যাট: {ext}")
        except Exception as e:
            logger.exception(f"ফাইল পড়তে সমস্যা: {path}")
            result.errors.append(f"পড়তে ব্যর্থ: {e}")

        return result

    # ------------------------------------------------------
    def _read_csv(self, path: Path, result: ExcelReadResult):
        """CSV ফাইল পড়ো — encoding নিজে শনাক্ত করে"""
        import chardet

        raw = path.read_bytes()[:100_000]
        detected = chardet.detect(raw)
        encoding = detected.get("encoding") or "utf-8"

        for enc in [encoding, "utf-8-sig", "utf-8", "cp1252", "latin-1"]:
            try:
                raw_df = pd.read_csv(
                    path, header=None, dtype=str, encoding=enc,
                    on_bad_lines="skip", engine="python",
                )
                break
            except Exception:
                continue
        else:
            result.errors.append("CSV পড়া যায়নি — encoding সমস্যা")
            return

        profile = self._analyze_sheet(raw_df, "Sheet1", False)
        result.sheets.append(profile)

        if profile.header_row is not None:
            df = self._build_dataframe(raw_df, profile)
            result.dataframes["Sheet1"] = df

    # ------------------------------------------------------
    def _read_excel(self, path: Path, result: ExcelReadResult):
        """Excel ফাইল পড়ো — সব Sheet, Hidden সহ"""
        hidden_sheets: set[str] = set()
        merged_map: dict[str, list[str]] = {}
        formula_sheets: set[str] = set()

        # openpyxl দিয়ে metadata নাও (শুধু .xlsx/.xlsm)
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
                for ws in wb.worksheets:
                    if ws.sheet_state != "visible":
                        hidden_sheets.add(ws.title)
                    merged_map[ws.title] = [str(r) for r in ws.merged_cells.ranges]
                    # Formula আছে কিনা দ্রুত চেক
                    for row in ws.iter_rows(max_row=min(ws.max_row or 1, 200)):
                        if any(isinstance(c.value, str) and c.value.startswith("=") for c in row):
                            formula_sheets.add(ws.title)
                            break
                wb.close()
            except Exception as e:
                result.warnings.append(f"Metadata পড়া যায়নি: {e}")

        # pandas দিয়ে ডেটা পড়ো (data_only — formula এর ফলাফল)
        engine_name = "xlrd" if path.suffix.lower() == ".xls" else "openpyxl"
        try:
            all_sheets = pd.read_excel(
                path, sheet_name=None, header=None, dtype=object, engine=engine_name
            )
        except Exception as e:
            result.errors.append(f"Excel পড়া যায়নি: {e}")
            return

        for sheet_name, raw_df in all_sheets.items():
            if raw_df.empty:
                continue

            profile = self._analyze_sheet(
                raw_df, sheet_name, sheet_name in hidden_sheets
            )
            profile.merged_ranges = merged_map.get(sheet_name, [])
            profile.merged_parsed = [
                pr for pr in (parse_merged_range(r) for r in profile.merged_ranges)
                if pr
            ]
            profile.has_formulas = sheet_name in formula_sheets
            result.sheets.append(profile)

            if profile.header_row is not None:
                df = self._build_dataframe(raw_df, profile)
                result.dataframes[sheet_name] = df
                if self.verbose:
                    logger.info(
                        f"📄 '{sheet_name}': হেডার সারি {profile.header_row + 1}, "
                        f"{len(df)} সারি ডেটা, {len(profile.detected_columns)} কলাম শনাক্ত "
                        f"({profile.likely_type})"
                    )
            else:
                result.warnings.append(f"'{sheet_name}' — হেডার খুঁজে পাওয়া যায়নি")

    # ------------------------------------------------------
    def _analyze_sheet(
        self, raw_df: pd.DataFrame, sheet_name: str, is_hidden: bool
    ) -> SheetProfile:
        """একটি Sheet বিশ্লেষণ করে হেডার ও কলাম শনাক্ত করো"""
        profile = SheetProfile(
            sheet_name=sheet_name,
            is_hidden=is_hidden,
            total_rows=len(raw_df),
            total_cols=len(raw_df.columns),
        )

        best_row, best_score, best_mapping, best_unmapped = None, 0, {}, []
        scan_limit = min(self.MAX_HEADER_SCAN_ROWS, len(raw_df))

        for idx in range(scan_limit):
            row = raw_df.iloc[idx]
            mapping: dict[str, str] = {}
            unmapped: list[str] = []
            filled = 0

            for col_idx, cell in enumerate(row):
                text = str(cell).strip() if cell is not None else ""
                if not text or text.lower() == "nan":
                    continue
                filled += 1
                canonical = match_column(text)
                if canonical and canonical not in mapping:
                    mapping[canonical] = text
                elif not canonical:
                    unmapped.append(text)

            # স্কোরিং: শনাক্ত কলাম বেশি + সারিতে অনেক cell ভরা
            score = len(mapping) * 10 + min(filled, 15)

            # হেডারের নিচে ডেটা আছে কিনা যাচাই
            if idx + 1 < len(raw_df):
                next_row = raw_df.iloc[idx + 1]
                next_filled = sum(
                    1 for c in next_row
                    if c is not None and str(c).strip() and str(c).lower() != "nan"
                )
                if next_filled < 2:
                    score -= 20

            if len(mapping) >= self.MIN_HEADER_MATCHES and score > best_score:
                best_row, best_score = idx, score
                best_mapping, best_unmapped = mapping, unmapped

        if best_row is None:
            return profile

        profile.header_row = best_row
        profile.data_start_row = best_row + 1
        profile.detected_columns = best_mapping
        profile.unmapped_columns = best_unmapped
        profile.confidence = min(best_score / 60.0, 1.0)
        profile.likely_type = self._guess_sheet_type(best_mapping, sheet_name)

        # Subtotal সারি খুঁজো
        for i in range(profile.data_start_row, len(raw_df)):
            if is_total_row(raw_df.iloc[i]):
                profile.subtotal_rows.append(i)

        profile.data_end_row = len(raw_df) - 1
        return profile

    # ------------------------------------------------------
    @staticmethod
    def _guess_sheet_type(mapping: dict[str, str], sheet_name: str) -> str:
        """Sheet টি কী ধরনের ডেটা ধারণ করে অনুমান করো"""
        name_lower = sheet_name.lower()
        keys = set(mapping.keys())

        if any(k in name_lower for k in [
            "entitle", "annexure", "annex", "এন্টাইটেল", "অনুমোদ", "প্রাপ্যতা", "কাঁচামাল"
        ]):
            return "entitlement"
        if any(k in name_lower for k in ["import", "im-4", "im4", "ais", "b/e", "আমদানি"]):
            return "import"
        if any(k in name_lower for k in ["consum", "খরচ", "ব্যবহার"]):
            return "consumption"
        if any(k in name_lower for k in ["export", "exp", "রপ্তানি"]):
            return "export"
        if any(k in name_lower for k in ["stock", "inventory", "মজুদ"]):
            return "inventory"

        # কলাম দেখে অনুমান — বিল নম্বর থাকলে নিশ্চিতভাবে আমদানি
        if "bill_number" in keys or "bill_date" in keys:
            return "import"
        if "entitled_quantity" in keys or "approved_quantity" in keys:
            return "entitlement"
        if {"hs_code", "quantity"}.issubset(keys):
            return "import"
        return "unknown"

    # ------------------------------------------------------
    def _build_dataframe(
        self, raw_df: pd.DataFrame, profile: SheetProfile
    ) -> pd.DataFrame:
        """হেডার প্রয়োগ করে পরিষ্কার DataFrame তৈরি করো"""
        header_vals = raw_df.iloc[profile.header_row]

        # কলাম নাম তৈরি — canonical নাম পেলে সেটা ব্যবহার করো
        columns: list[str] = []
        seen: dict[str, int] = {}
        for i, cell in enumerate(header_vals):
            text = str(cell).strip() if cell is not None else ""
            if not text or text.lower() == "nan":
                name = f"col_{i}"
            else:
                canonical = match_column(text)
                name = canonical if canonical else normalize_text(text).replace(" ", "_")
                if not name:
                    name = f"col_{i}"
            # ডুপ্লিকেট নাম হ্যান্ডল
            if name in seen:
                seen[name] += 1
                name = f"{name}_{seen[name]}"
            else:
                seen[name] = 0
            columns.append(name)

        profile.column_positions = {name: i for i, name in enumerate(columns)}

        df = raw_df.iloc[profile.data_start_row:].copy()
        df.columns = columns
        # মূল Excel সারি নম্বর ধরে রাখো (merged cell মেলাতে প্রয়োজন)
        df["_raw_row"] = list(range(profile.data_start_row, len(raw_df)))
        df = df.reset_index(drop=True)

        # Merged Cell — উপরের মান নিচে ছড়িয়ে দাও (forward fill)
        # শুধু text কলামে, সংখ্যায় নয়
        text_cols = ["item_name", "hs_code", "bill_number", "bill_date", "supplier"]
        for col in text_cols:
            if col in df.columns:
                df[col] = df[col].replace("", np.nan).ffill()

        # সম্পূর্ণ খালি সারি বাদ দাও
        df = df.dropna(how="all")

        # Subtotal সারি চিহ্নিত করো (বাদ দেওয়ার জন্য)
        df["_is_subtotal"] = df.apply(is_total_row, axis=1)

        # সংখ্যা কলাম পরিষ্কার করো
        numeric_cols = [
            "quantity", "approved_quantity", "entitled_quantity",
            "closing_stock", "opening_stock", "bonding_capacity",
            "capacity_100", "capacity_80", "capacity_60", "warehouse_capacity",
            "local_purchase", "unit_price", "value_usd",
            "value_bdt", "duty_rate", "vat_rate", "at_rate", "rd_rate",
            "sd_rate", "ait_rate", "duty_paid", "vat_paid", "rd_paid",
            "sd_paid", "at_paid", "ait_paid", "total_tax", "exchange_rate",
        ]
        for col in numeric_cols:
            if col in df.columns:
                df[col] = df[col].apply(clean_number)

        # HS কোড পরিষ্কার করো
        if "hs_code" in df.columns:
            df["hs_code"] = df["hs_code"].apply(clean_hs_code)

        # তারিখ কলাম
        for col in ["bill_date", "lc_date", "bl_date"]:
            if col in df.columns:
                df[col] = pd.to_datetime(df[col], errors="coerce", dayfirst=True)

        # টেক্সট কলাম পরিষ্কার
        for col in ["item_name", "commercial_name", "supplier", "unit"]:
            if col in df.columns:
                df[col] = df[col].apply(
                    lambda x: str(x).strip() if x is not None and str(x).lower() != "nan" else None
                )

        return df

    # ------------------------------------------------------
    def get_clean_data(
        self, df: pd.DataFrame, drop_subtotals: bool = True
    ) -> pd.DataFrame:
        """বিশ্লেষণের জন্য পরিষ্কার ডেটা — subtotal বাদ"""
        out = df.copy()
        if drop_subtotals and "_is_subtotal" in out.columns:
            out = out[~out["_is_subtotal"]]
            out = out.drop(columns=["_is_subtotal"])
        return out.reset_index(drop=True)

    # ------------------------------------------------------
    def preview(self, file_path: str | Path, rows: int = 10) -> dict:
        """UI তে দেখানোর জন্য ফাইলের প্রিভিউ"""
        result = self.read(file_path)
        preview_data = {
            "file_name": result.file_name,
            "sheets": [],
            "errors": result.errors,
            "warnings": result.warnings,
        }

        for profile in result.sheets:
            df = result.dataframes.get(profile.sheet_name)
            sheet_info = {
                "name": profile.sheet_name,
                "is_hidden": profile.is_hidden,
                "header_row": (profile.header_row + 1) if profile.header_row is not None else None,
                "total_rows": profile.total_rows,
                "data_rows": len(df) if df is not None else 0,
                "detected_columns": profile.detected_columns,
                "unmapped_columns": profile.unmapped_columns[:20],
                "subtotal_rows": len(profile.subtotal_rows),
                "has_formulas": profile.has_formulas,
                "merged_cells": len(profile.merged_ranges),
                "confidence": round(profile.confidence, 2),
                "likely_type": profile.likely_type,
                "sample": [],
            }
            if df is not None and not df.empty:
                sample_df = self.get_clean_data(df).head(rows)
                sheet_info["sample"] = sample_df.astype(str).to_dict("records")
            preview_data["sheets"].append(sheet_info)

        return preview_data


# ==========================================================
# Convenience Functions
# ==========================================================

def read_excel(file_path: str | Path) -> ExcelReadResult:
    """দ্রুত ফাইল পড়ার শর্টকাট"""
    return ExcelEngine().read(file_path)


def preview_excel(file_path: str | Path, rows: int = 10) -> dict:
    """দ্রুত প্রিভিউ"""
    return ExcelEngine().preview(file_path, rows)
