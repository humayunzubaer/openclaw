"""
Bond Register Template — তফসিল-১ অনুযায়ী বন্ড রেজিস্টার
============================================================

আইনি ভিত্তি:
    সরাসরি ও প্রচ্ছন্ন রপ্তানিমুখী (পোশাক শিল্প ব্যতীত) শিল্প প্রতিষ্ঠান
    (ওয়্যারহাউস পদ্ধতির আওতায় সাময়িক আমদানি, ওয়্যারহাউস পরিচালনা ও
    কার্যপদ্ধতি) বিধিমালা, ২০২৪ [এসআরও ২১২-আইন/২০২৪/৬৪/কাস্টমস]
    — বিধি ৭ ও তফসিল-১

উদ্দেশ্য:
    এককালীন বন্ডিং ক্যাপাসিটি যাচাইয়ের জন্য ইন্টু-বন্ড ও এক্স-বন্ড একত্রে
    থাকা কনজাম্পশন রেজিস্টার আবশ্যক। ইহা ব্যতীত কোনো মুহূর্তের প্রকৃত
    মজুত নির্ণয় সম্ভব নহে।

    এই মডিউল —
      ১) নিরীক্ষককে তফসিল-১ এর ছক প্রদর্শন করে
      ২) পূরণযোগ্য Excel টেমপ্লেট ডাউনলোডের সুযোগ দেয়
      ৩) নিরীক্ষক এই প্রশ্ন এড়াইয়া (skip) যাইতে পারেন — সেই ক্ষেত্রে
         বন্ডিং ক্যাপাসিটি সংক্রান্ত ফাইন্ডিংস প্রদান করা হইবে না
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils.logger import logger


# ==========================================================
# তফসিল-১ এর হুবহু ১৬ কলাম
# ==========================================================
# (কলাম নং, শিরোনাম, প্রস্থ, ধরন)

SCHEDULE_1_COLUMNS: list[tuple[str, str, int, str]] = [
    ("১",  "বিল অব এন্ট্রি নম্বর ও তারিখ", 22, "text"),
    ("২",  "আমদানি কাস্টম হাউস/স্টেশনের নাম", 22, "text"),
    ("৩",  "পণ্যচালান কাস্টম হাউস/স্টেশন হইতে ছাড়করণের তারিখ\n(এক্সিট নোটের তারিখ)", 22, "date"),
    ("৪",  "এলসি নং ও তারিখ", 20, "text"),
    ("৫",  "এইচএস কোড", 14, "text"),
    ("৬",  "পণ্যের বাণিজ্যিক বর্ণনা", 34, "text"),
    ("৭",  "কাঁচামালের পরিমাণ\n— কেজি", 14, "number"),
    ("৮",  "কাঁচামালের পরিমাণ\n— মিটার", 14, "number"),
    ("৯",  "কাঁচামালের পরিমাণ\n— গজ", 14, "number"),
    ("১০", "কাঁচামালের মোট মূল্য\n(ইউএস ডলার/অন্যান্য মুদ্রা)", 18, "number"),
    ("১১", "ইন্টু বন্ডের তারিখ", 15, "date"),
    ("১২", "বন্ড অফিসারের স্বাক্ষর", 16, "text"),
    ("১৩", "এক্স বন্ডের তারিখ", 15, "date"),
    ("১৪", "এক্সবন্ডকৃত পণ্যের পরিমাণ", 18, "number"),
    ("১৫", "সমাপনী মজুদ", 15, "number"),
    ("১৬", "বন্ড অফিসারের স্বাক্ষর", 16, "text"),
]

# হেডার তথ্যক্ষেত্র (তফসিল-১ এর ১–৬ নম্বর)
SCHEDULE_1_HEADER_FIELDS: list[str] = [
    "১। শিল্প প্রতিষ্ঠানের নাম ও ব্যবসা শনাক্তকরণ নম্বর (BIN)",
    "২। ঠিকানা",
    "৩। বন্ড রেজিস্ট্রার নম্বর",
    "৪। ওয়্যারহাউস লাইসেন্স নম্বর",
    "৫। বন্ড রেজিস্ট্রারের পৃষ্ঠা সংখ্যা",
    "৬। ব্যবস্থাপনা কর্তৃপক্ষের পক্ষে স্বাক্ষর",
]


# ==========================================================
# নিরীক্ষককে প্রদর্শনের জন্য অনুরোধ-বার্তা
# ==========================================================

REGISTER_REQUEST_TITLE = "কনজাম্পশন রেজিস্টার (তফসিল-১) প্রয়োজন"

REGISTER_REQUEST_MESSAGE = """
এককালীন বন্ডিং ক্যাপাসিটি যাচাই করিতে হইলে নিরীক্ষাধীন মেয়াদের যে কোনো
মুহূর্তে ওয়্যারহাউসে প্রকৃতপক্ষে কী পরিমাণ কাঁচামাল মজুত ছিল তাহা জানা আবশ্যক।

কারণ — এককালীন বন্ডিং ক্যাপাসিটি লঙ্ঘন নির্ণীত হয় নিম্নরূপে:

    মজুত(t) = প্রারম্ভিক জের + ইন্টু-বন্ড (t পর্যন্ত) − এক্স-বন্ড (t পর্যন্ত)

শুধুমাত্র আমদানির তথ্য (MIS) দিয়া উত্তোলন (এক্স-বন্ড) বাদ দেওয়া সম্ভব নহে;
ফলে মজুত ক্রমবর্ধমান দেখাইবে এবং লঙ্ঘন অতিরঞ্জিত হইবে।

★ অনুগ্রহ করিয়া ইন্টু-বন্ড ও এক্স-বন্ড একত্রে সন্নিবেশিত কনজাম্পশন
  রেজিস্টার (তফসিল-১ অনুযায়ী বন্ড রেজিস্ট্রার) প্রাপ্যতা শীটের সহিত
  আপলোড করুন।

আইনি ভিত্তি: এসআরও ২১২-আইন/২০২৪, বিধি ৭ — "প্রত্যেক ওয়্যারহাউস
লাইসেন্সধারী প্রতিষ্ঠানকে তফসিল-১ অনুযায়ী ২ (দুই) কপি করিয়া বন্ড
রেজিস্ট্রার সংরক্ষণ করিতে হইবে।"

আপনি চাহিলে —
  • নিচের বাটন হইতে পূরণযোগ্য টেমপ্লেট ডাউনলোড করিতে পারেন
  • অথবা এই ধাপটি এড়াইয়া (Skip) যাইতে পারেন

⚠ এড়াইয়া গেলে বন্ডিং ক্যাপাসিটি সংক্রান্ত ফাইন্ডিংস ও দাবি প্রদান করা
  হইবে না। রেজিস্টার পরবর্তীতে পাওয়া গেলে পুনরায় যাচাই করা যাইবে।
""".strip()

SKIP_NOTE = (
    "কনজাম্পশন রেজিস্টার (তফসিল-১) পাওয়া না যাওয়ায় এককালীন বন্ডিং "
    "ক্যাপাসিটি সংক্রান্ত যাচাই সম্পন্ন করা সম্ভব হয় নাই। ইন্টু-বন্ড ও "
    "এক্স-বন্ড একত্রে সন্নিবেশিত রেজিস্টার ব্যতীত নিরীক্ষাধীন মেয়াদের "
    "কোনো মুহূর্তের প্রকৃত মজুত নির্ণয় করা যায় না বিধায় এই বিষয়ে কোনো "
    "ফাইন্ডিংস বা দাবি উত্থাপন করা হইল না। প্রতিষ্ঠান কর্তৃক উক্ত রেজিস্টার "
    "উপস্থাপিত হইলে বিষয়টি পুনরায় যাচাই করা হইবে।"
)


# ==========================================================
# টেমপ্লেট জেনারেটর
# ==========================================================

# নকশা ধ্রুবক
_FONT = "Arial"
_C_HEAD = "1F3864"
_C_SUB = "D9E2F3"
_C_INSTR = "FFF2CC"
_THIN = Side(style="thin", color="8EA9DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def create_bond_register_template(
    out_path: str | Path,
    company_name: str = "",
    bond_license: str = "",
    period_from: Optional[date] = None,
    period_to: Optional[date] = None,
    sample_rows: int = 200,
) -> Path:
    """
    তফসিল-১ অনুযায়ী পূরণযোগ্য বন্ড রেজিস্টার টেমপ্লেট তৈরি করো।

    নিরীক্ষক ইহা ডাউনলোড করিয়া প্রতিষ্ঠানকে পূরণ করিতে দিবেন।
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # ============ শীট ১: বন্ড রেজিস্টার ============
    ws = wb.active
    ws.title = "বন্ড রেজিস্টার"

    ncols = len(SCHEDULE_1_COLUMNS)

    # --- শিরোনাম ---
    titles = [
        "গণপ্রজাতন্ত্রী বাংলাদেশ সরকার",
        "জাতীয় রাজস্ব বোর্ড",
        "বন্ড রেজিস্ট্রার",
        "[তফসিল-১] [বিধি ৭ দ্রষ্টব্য] — এসআরও ২১২-আইন/২০২৪/৬৪/কাস্টমস",
    ]
    for i, t in enumerate(titles, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(
            name=_FONT, size=13 if i == 3 else (11 if i < 3 else 9),
            bold=i <= 3, italic=i == 4,
            color="FFFFFF" if i == 3 else "000000",
        )
        if i == 3:
            c.fill = PatternFill("solid", fgColor=_C_HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ncols)
    ws.row_dimensions[3].height = 24

    # --- হেডার তথ্যক্ষেত্র ---
    r = 6
    for label in SCHEDULE_1_HEADER_FIELDS:
        c = ws.cell(row=r, column=1, value=label)
        c.font = Font(name=_FONT, size=10, bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)

        # পূরণযোগ্য ঘর
        v = ws.cell(row=r, column=5, value="")
        v.fill = PatternFill("solid", fgColor="FFFFCC")
        v.border = _BORDER
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
        r += 1

    # প্রি-ফিল
    if company_name:
        ws.cell(row=6, column=5, value=company_name)
    if bond_license:
        ws.cell(row=9, column=5, value=bond_license)

    # --- মেয়াদ ---
    r += 1
    if period_from and period_to:
        p = (f"নিরীক্ষাধীন মেয়াদ: {period_from.strftime('%d.%m.%Y')} "
             f"হইতে {period_to.strftime('%d.%m.%Y')}")
        c = ws.cell(row=r, column=1, value=p)
        c.font = Font(name=_FONT, size=11, bold=True, color="C00000")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        r += 2

    # --- নির্দেশনা ---
    instr = (
        "★ পূরণের নির্দেশনা: প্রতিটি ইন্টু-বন্ড ও এক্স-বন্ড লেনদেন "
        "তারিখক্রমে লিপিবদ্ধ করুন। কলাম ৭/৮/৯-এ প্রযোজ্য এককে পরিমাণ "
        "লিখুন (কেজি অবশ্যই পূরণ করুন — এককালীন বন্ডিং ক্যাপাসিটি "
        "যাচাইয়ের জন্য আবশ্যক)। কলাম ১৫-এ প্রতিটি লেনদেনের পর সমাপনী "
        "মজুদ লিখুন। হলুদ ঘরসমূহ পূরণযোগ্য।"
    )
    c = ws.cell(row=r, column=1, value=instr)
    c.font = Font(name=_FONT, size=9, italic=True)
    c.fill = PatternFill("solid", fgColor=_C_INSTR)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    ws.row_dimensions[r].height = 42
    r += 2

    # --- কলাম শিরোনাম ---
    header_row = r
    for i, (num, title, width, _) in enumerate(SCHEDULE_1_COLUMNS, start=1):
        c = ws.cell(row=header_row, column=i, value=title)
        c.font = Font(name=_FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_C_HEAD)
        c.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        c.border = _BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[header_row].height = 58

    # --- কলাম নম্বরের সারি ---
    num_row = header_row + 1
    for i, (num, _, _, _) in enumerate(SCHEDULE_1_COLUMNS, start=1):
        c = ws.cell(row=num_row, column=i, value=f"({num})")
        c.font = Font(name=_FONT, size=8, bold=True)
        c.fill = PatternFill("solid", fgColor=_C_SUB)
        c.alignment = Alignment(horizontal="center")
        c.border = _BORDER

    # --- পূরণযোগ্য খালি সারি ---
    data_start = num_row + 1
    for row in range(data_start, data_start + sample_rows):
        for i, (_, _, _, kind) in enumerate(SCHEDULE_1_COLUMNS, start=1):
            c = ws.cell(row=row, column=i)
            c.border = _BORDER
            c.font = Font(name=_FONT, size=9)
            if kind == "number":
                c.number_format = '#,##0.000;(#,##0.000);-'
                c.alignment = Alignment(horizontal="right")
            elif kind == "date":
                c.number_format = 'DD.MM.YYYY'
                c.alignment = Alignment(horizontal="center")

    ws.freeze_panes = ws.cell(row=data_start, column=1)

    # ============ শীট ২: নির্দেশনা ও আইনি ভিত্তি ============
    ws2 = wb.create_sheet("নির্দেশনা")
    ws2.column_dimensions["A"].width = 4
    ws2.column_dimensions["B"].width = 110
    ws2.sheet_view.showGridLines = False

    c = ws2.cell(row=1, column=2, value="বন্ড রেজিস্টার পূরণ ও আপলোড নির্দেশনা")
    c.font = Font(name=_FONT, size=14, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=_C_HEAD)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws2.merge_cells("B1:B1")
    ws2.row_dimensions[1].height = 28

    rr = 3
    for para in REGISTER_REQUEST_MESSAGE.split("\n\n"):
        c = ws2.cell(row=rr, column=2, value=para.strip())
        c.font = Font(name=_FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        lines = max(2, len(para) // 95 + para.count("\n") + 1)
        ws2.row_dimensions[rr].height = lines * 15
        rr += 1

    rr += 1
    c = ws2.cell(row=rr, column=2, value="কলামভিত্তিক নির্দেশনা")
    c.font = Font(name=_FONT, size=12, bold=True, color="1F3864")
    rr += 1
    col_notes = [
        "কলাম ১–৬ : আমদানি/সংগ্রহ সংক্রান্ত তথ্য — বিল অব এন্ট্রি, এলসি, "
        "এইচএস কোড ও পণ্যের বাণিজ্যিক বর্ণনা",
        "কলাম ৭–৯ : পরিমাণ — কেজি, মিটার, গজ। ★ কেজি কলাম অবশ্যই পূরণ করুন; "
        "সকল কাঁচামালের সমষ্টি একই এককে না আনিলে এককালীন বন্ডিং ক্যাপাসিটি "
        "(মেট্রিক টন) এর সহিত তুলনা করা সম্ভব নহে",
        "কলাম ১১ : ইন্টু বন্ডের তারিখ — ওয়্যারহাউসে প্রবেশের তারিখ। "
        "ছাড়করণের ৫ দিনের মধ্যে হইতে হইবে [এসআরও ২১২/২০২৪, বিধি ৮(১)]",
        "কলাম ১৩–১৪ : এক্স বন্ড — উত্তোলনের তারিখ ও পরিমাণ। ইহাই মজুত "
        "হইতে বিয়োজিত হইবে",
        "কলাম ১৫ : সমাপনী মজুদ — প্রতিটি লেনদেনের পরের স্থিতি",
        "স্থানীয় উৎস হইতে সংগৃহীত কাঁচামালও অনুরূপভাবে ইন্টু বন্ড করিতে "
        "হইবে [বিধি ৮(২)] — তাই সেগুলিও এই রেজিস্টারে অন্তর্ভুক্ত করুন",
    ]
    for note in col_notes:
        ws2.cell(row=rr, column=1, value="▸").font = Font(name=_FONT, size=10)
        c = ws2.cell(row=rr, column=2, value=note)
        c.font = Font(name=_FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.row_dimensions[rr].height = max(30, len(note) // 95 * 15 + 15)
        rr += 1

    wb.save(out_path)
    logger.success(f"তফসিল-১ বন্ড রেজিস্টার টেমপ্লেট তৈরি হইয়াছে: {out_path}")
    return out_path


# ==========================================================
# নিরীক্ষকের সিদ্ধান্ত ধারণকারী কাঠামো
# ==========================================================

@dataclass
class RegisterDecision:
    """
    কনজাম্পশন রেজিস্টার সংক্রান্ত নিরীক্ষকের সিদ্ধান্ত।

    UI-তে তিনটি বিকল্প দেখানো হইবে:
        ১) রেজিস্টার আপলোড করুন
        ২) টেমপ্লেট ডাউনলোড করুন
        ৩) এই ধাপ এড়াইয়া যান (Skip)
    """
    provided: bool = False               # রেজিস্টার পাওয়া গিয়াছে কিনা
    skipped: bool = False                # নিরীক্ষক এড়াইয়া গিয়াছেন কিনা
    file_path: str = ""
    template_downloaded: bool = False
    skip_reason: str = ""

    @property
    def can_check_capacity(self) -> bool:
        """বন্ডিং ক্যাপাসিটি যাচাই সম্ভব কিনা"""
        return self.provided and not self.skipped

    @property
    def note(self) -> str:
        """প্রতিবেদনে বসাইবার টীকা"""
        if self.can_check_capacity:
            return (
                "কনজাম্পশন রেজিস্টার (তফসিল-১) পাওয়া গিয়াছে; ইন্টু-বন্ড ও "
                "এক্স-বন্ড তথ্যের ভিত্তিতে প্রতিটি মুহূর্তের প্রকৃত মজুত "
                "নির্ণয়পূর্বক এককালীন বন্ডিং ক্যাপাসিটি যাচাই করা হইয়াছে।"
            )
        if self.skipped:
            return SKIP_NOTE + (
                f" [নিরীক্ষকের মন্তব্য: {self.skip_reason}]"
                if self.skip_reason else ""
            )
        return SKIP_NOTE


def get_register_prompt(
    company_name: str = "", period_label: str = ""
) -> dict:
    """
    UI-তে দেখানোর জন্য অনুরোধ-বার্তা ও বিকল্পসমূহ।
    FastAPI/Electron ইহা ব্যবহার করিবে।
    """
    return {
        "title": REGISTER_REQUEST_TITLE,
        "message": REGISTER_REQUEST_MESSAGE,
        "company_name": company_name,
        "period": period_label,
        "schedule_columns": [
            {"no": num, "title": title.replace("\n", " ")}
            for num, title, _, _ in SCHEDULE_1_COLUMNS
        ],
        "legal_basis": (
            "এসআরও ২১২-আইন/২০২৪/৬৪/কাস্টমস — বিধি ৭ ও তফসিল-১"
        ),
        "actions": [
            {
                "id": "upload",
                "label": "রেজিস্টার আপলোড করুন",
                "primary": True,
                "description": "পূরণকৃত তফসিল-১ ফাইল নির্বাচন করুন",
            },
            {
                "id": "download_template",
                "label": "টেমপ্লেট ডাউনলোড করুন",
                "primary": False,
                "description": "তফসিল-১ অনুযায়ী পূরণযোগ্য Excel ছক",
            },
            {
                "id": "skip",
                "label": "এই ধাপ এড়াইয়া যান",
                "primary": False,
                "warning": (
                    "বন্ডিং ক্যাপাসিটি সংক্রান্ত ফাইন্ডিংস ও দাবি "
                    "প্রদান করা হইবে না"
                ),
            },
        ],
    }


__all__ = [
    "SCHEDULE_1_COLUMNS", "SCHEDULE_1_HEADER_FIELDS",
    "REGISTER_REQUEST_TITLE", "REGISTER_REQUEST_MESSAGE", "SKIP_NOTE",
    "create_bond_register_template", "RegisterDecision", "get_register_prompt",
]
