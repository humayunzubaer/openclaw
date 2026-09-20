"""
Word প্রতিবেদন — Nikosh (ইউনিকোড) ও SutonnyMJ (বিজয়) উভয় রূপে
==================================================================

ফন্ট-নীতি
---------
    Nikosh      → ইউনিকোড। লেখা যেমন আছে তেমনই বসে। ঝুঁকি নাই।
    SutonnyMJ   → বিজয় ASCII। বাংলা অংশ রূপান্তরিত হইয়া বসে;
                  ইংরেজি ও সংখ্যা **আলাদা রানে** ইংরেজি ফন্টে বসে।

★ কেন আলাদা রান
    বিজয়ে 'K' মানে 'ক'। তাই SutonnyMJ রানে ইংরেজি রাখিলে তাহা বাংলা
    আবর্জনা হইয়া দেখায়। `bijoy.split_runs()` ইহাই ঠেকায়।

Word-এ বাংলা বসাইবার সূক্ষ্মতা
------------------------------
    Word বাংলাকে "complex script" ধরে। কাজেই কেবল `w:ascii` দিলে হয় না —
    `w:cs` (complex script ফন্ট) ও `w:szCs` (complex script আকার) দুইটিই
    বসাইতে হয়। নতুবা Word নিজের পছন্দের ফন্টে ফেলিয়া দেয়।
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

from services.bijoy import split_runs, unicode_to_bijoy
from utils.logger import logger


# ==========================================================
# ফন্ট-বিন্যাস
# ==========================================================

@dataclass(frozen=True)
class FontScheme:
    """প্রতিবেদনের ফন্ট-নীতি"""
    key: str
    label: str
    bangla_font: str          # বাংলার ফন্টের নাম
    latin_font: str           # ইংরেজি/সংখ্যার ফন্ট
    convert_to_bijoy: bool    # বাংলা কি বিজয় ASCII-তে বদলাইবে?
    base_size: float          # পয়েন্ট
    note: str = ""


NIKOSH = FontScheme(
    key="nikosh",
    label="Nikosh (ইউনিকোড)",
    bangla_font="Nikosh",
    latin_font="Times New Roman",
    convert_to_bijoy=False,
    base_size=12.0,
    note=("ইউনিকোড — লেখা অপরিবর্তিত থাকে। যেকোনো কম্পিউটারে খোলা যায়; "
          "খোঁজা ও অনুলিপি করা চলে।"),
)

SUTONNY = FontScheme(
    key="sutonnymj",
    label="SutonnyMJ (বিজয়)",
    bangla_font="SutonnyMJ",
    latin_font="Times New Roman",
    convert_to_bijoy=True,
    base_size=14.0,           # বিজয় ফন্ট দৃশ্যত ছোট — তাই বড় আকার
    note=("বিজয় ASCII — বাংলা রূপান্তরিত হইয়া বসে। যে কম্পিউটারে "
          "SutonnyMJ নাই সেখানে বিকৃত দেখাইবে।"),
)

SCHEMES: dict[str, FontScheme] = {s.key: s for s in (NIKOSH, SUTONNY)}


def scheme_for(name: str | None) -> FontScheme:
    return SCHEMES.get((name or "nikosh").strip().lower(), NIKOSH)


# ==========================================================
# নিম্নস্তর — রান বসানো
# ==========================================================

def _clear(paragraph) -> None:
    """অনুচ্ছেদের বিদ্যমান রান মুছে — খালি ফন্টবিহীন রান রাখে না"""
    for r in list(paragraph.runs):
        r._element.getparent().remove(r._element)


def _apply_font(run, font_name: str, size_pt: float,
                bold: bool = False, italic: bool = False) -> None:
    """একটি রানে ফন্ট বসায় — complex-script ক্ষেত্রসহ"""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), font_name)

    # complex-script আকার — না দিলে Word বাংলাকে অন্য মাপে ফেলে
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = rpr.makeelement(qn("w:szCs"), {})
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(size_pt * 2)))


def write_text(paragraph, text: str, scheme: FontScheme,
               size_pt: Optional[float] = None,
               bold: bool = False, italic: bool = False) -> None:
    """
    একটি অনুচ্ছেদে লেখা বসায় — ফন্ট-নীতি মানিয়া।

    SutonnyMJ হইলে বাংলা ও অ-বাংলা অংশ পৃথক রানে যায়, যাহাতে ইংরেজি
    কখনো বিজয়-গ্লিফ হইয়া না দেখায়।
    """
    size = size_pt or scheme.base_size
    if not text:
        return

    if not scheme.convert_to_bijoy:
        run = paragraph.add_run(text)
        _apply_font(run, scheme.bangla_font, size, bold, italic)
        return

    for kind, seg in split_runs(text):
        if not seg:
            continue
        if kind == "bn":
            run = paragraph.add_run(unicode_to_bijoy(seg))
            _apply_font(run, scheme.bangla_font, size, bold, italic)
        else:
            run = paragraph.add_run(seg)
            _apply_font(run, scheme.latin_font, size - 2, bold, italic)


# ==========================================================
# প্রতিবেদন-নির্মাতা
# ==========================================================

class ReportDocument:
    """বাংলা নিরীক্ষা প্রতিবেদনের Word নথি"""

    def __init__(self, scheme: FontScheme | str = NIKOSH):
        self.scheme = scheme if isinstance(scheme, FontScheme) else scheme_for(scheme)
        self.doc = Document()
        self._setup_page()

    # ------------------------------------------------------
    def _setup_page(self) -> None:
        for s in self.doc.sections:
            s.top_margin = Cm(2.5)
            s.bottom_margin = Cm(2.5)
            s.left_margin = Cm(3.0)      # বাঁ পাশে বাঁধাইয়ের জায়গা
            s.right_margin = Cm(2.0)
        style = self.doc.styles["Normal"]
        style.font.name = self.scheme.bangla_font
        style.font.size = Pt(self.scheme.base_size)
        pf = style.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.4

    # ------------------------------------------------------
    def title(self, text: str, sub: str = "") -> "ReportDocument":
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        write_text(p, text, self.scheme,
                   size_pt=self.scheme.base_size + 3, bold=True)
        if sub:
            q = self.doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            write_text(q, sub, self.scheme, size_pt=self.scheme.base_size - 1)
        return self

    def heading(self, text: str) -> "ReportDocument":
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        write_text(p, text, self.scheme,
                   size_pt=self.scheme.base_size + 1, bold=True)
        return self

    def para(self, text: str, justify: bool = True,
             indent_cm: float = 0.0) -> "ReportDocument":
        p = self.doc.add_paragraph()
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent_cm:
            p.paragraph_format.left_indent = Cm(indent_cm)
        write_text(p, text, self.scheme)
        return self

    def numbered(self, items: Iterable[str], start: int = 1
                 ) -> "ReportDocument":
        """বাংলা অঙ্কে ক্রমিক — Word-এর তালিকা নহে, যাহাতে ফন্ট ঠিক থাকে"""
        from knowledge.report_style import _to_bangla_digits as bn
        for i, it in enumerate(items, start):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-1.0)
            write_text(p, f"{bn(str(i))}। {it}", self.scheme)
        return self

    def table(self, headers: list[str], rows: list[list[str]]
              ) -> "ReportDocument":
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for c, h in enumerate(headers):
            cell = t.rows[0].cells[c]
            _clear(cell.paragraphs[0])
            write_text(cell.paragraphs[0], h, self.scheme,
                       size_pt=self.scheme.base_size - 1, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for c, v in enumerate(row[:len(headers)]):
                _clear(cells[c].paragraphs[0])
                write_text(cells[c].paragraphs[0], str(v), self.scheme,
                           size_pt=self.scheme.base_size - 1)
        return self

    def page_break(self) -> "ReportDocument":
        self.doc.add_page_break()
        return self

    def spacer(self) -> "ReportDocument":
        self.doc.add_paragraph()
        return self

    # ------------------------------------------------------
    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        logger.info(f"প্রতিবেদন লেখা হইল ({self.scheme.label}): {path}")
        return path


__all__ = ["ReportDocument", "FontScheme", "NIKOSH", "SUTONNY",
           "SCHEMES", "scheme_for", "write_text"]
