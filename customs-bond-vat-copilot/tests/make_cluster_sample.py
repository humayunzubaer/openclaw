"""
ক্লাস্টার-সহ বাস্তবসম্মত প্রাপ্যতা শীট
=======================================
ক্লাস্টার-১ : ৪টি কাপড় একত্রে প্রাপ্যতা ৩,৭০,০০০ YDS (Merged Cell)
ক্লাস্টার-২ : ৩টি আনুষঙ্গিক একত্রে প্রাপ্যতা ৩৫,০০,০০০ PCS (খালি সারি পদ্ধতি)
একক আইটেম  : বাকিগুলো
"""
import openpyxl
from openpyxl.styles import Font, Alignment
from pathlib import Path

OUT = Path(__file__).parent / "sample_cluster.xlsx"
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "প্রাপ্যতা"

ws["A1"] = "গণপ্রজাতন্ত্রী বাংলাদেশ সরকার | কাস্টমস বন্ড কমিশনারেট, ঢাকা"
ws["A2"] = "প্রতিষ্ঠান: ABC Apparels Ltd. | বন্ড লাইসেন্স নং: DHK/BOND/1234"
ws["A3"] = "এককালীন বন্ডিং ক্যাপাসিটি: টাকা ৪,৫০,০০,০০০/-"
ws["A4"] = "নিরীক্ষা মেয়াদ: ০১.০১.২০২৪ হইতে ৩১.১২.২০২৪ (এই নিরীক্ষা সম্পন্ন শেষে প্রাপ্যতা প্রদান করা হইল)"

PERIOD = "০১.০১.২০২৫ থেকে ৩১.১২.২০২৫ পর্যন্ত সময়ের জন্য প্রদত্ত/প্রস্তাবিত আমদানি প্রাপ্যতা"
headers = ["ক্রমিক নং", "এইচ.এস কোড", "কাঁচামালের বিবরণ", "একক", PERIOD,
           "সমাপনী মজুত", "এককালীন বন্ডিং ক্যাপাসিটি", "একক মূল্য (USD)",
           "CD %", "RD %", "SD %", "VAT %", "AT %", "AIT %"]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=6, column=i, value=h)
    c.font = Font(bold=True, size=9)
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[6].height = 60
for col in "EFG":
    ws.column_dimensions[col].width = 20
ws.column_dimensions["C"].width = 40

r = 7

# ===== ক্লাস্টার ১ : কাপড় (৪টি) — Merged Cell পদ্ধতি =====
CLUSTER1 = [
    ("5208.11.00", "Woven Fabrics of Cotton, Unbleached"),
    ("5209.11.00", "Woven Fabrics of Cotton, >200g/m2"),
    ("5407.61.00", "Woven Fabrics of Polyester Filament"),
    ("5513.11.00", "Woven Fabrics of Polyester Staple Fibre"),
]
c1_start = r
ws.cell(row=r, column=1, value="১")
for hs, name in CLUSTER1:
    ws.cell(row=r, column=2, value=hs)
    ws.cell(row=r, column=3, value=name)
    r += 1
c1_end = r - 1
# একত্রে প্রাপ্যতা — ৪ সারি জুড়ে একটিই ঘর
for col, val in [(4, "YDS"), (5, 370000), (6, 35000), (7, 150000), (8, 1.50),
                 (9, 5), (10, 0), (11, 0), (12, 15), (13, 5), (14, 5)]:
    ws.cell(row=c1_start, column=col, value=val)
    ws.merge_cells(start_row=c1_start, start_column=col, end_row=c1_end, end_column=col)

# ===== ক্লাস্টার ২ : আনুষঙ্গিক (৩টি) — খালি সারি পদ্ধতি =====
CLUSTER2 = [
    ("9606.21.00", "Buttons of Plastics"),
    ("9607.11.00", "Slide Fasteners with Metal Teeth"),
    ("5807.10.00", "Woven Labels and Badges"),
]
ws.cell(row=r, column=1, value="২")
c2_start = r
for i, (hs, name) in enumerate(CLUSTER2):
    ws.cell(row=r, column=2, value=hs)
    ws.cell(row=r, column=3, value=name)
    ws.cell(row=r, column=4, value="PCS")
    if i == len(CLUSTER2) - 1:          # শেষ সারিতে একত্রে প্রাপ্যতা
        ws.cell(row=r, column=5, value=3500000)
        ws.cell(row=r, column=6, value=400000)
        ws.cell(row=r, column=7, value=1800000)
        ws.cell(row=r, column=8, value=0.020)
        for col, v in [(9,10),(10,3),(11,0),(12,15),(13,5),(14,5)]:
            ws.cell(row=r, column=col, value=v)
    r += 1

# ===== একক কাঁচামাল =====
SINGLES = [
 ("৩","5806.32.00","Narrow Woven Fabrics of Man-made Fibre","MTR",80000,8000,40000,0.15,10,3,0,15,5,5),
 ("৪","4819.10.00","Cartons, Boxes of Corrugated Paper","PCS",65000,5000,30000,0.45,15,3,0,15,5,5),
 ("৫","5401.10.00","Sewing Thread of Synthetic Filaments","KG",18000,2000,9000,3.10,5,0,0,15,5,5),
]
for row in SINGLES:
    for i, v in enumerate(row, start=1):
        ws.cell(row=r, column=i, value=v)
    r += 1

ws.cell(row=r, column=3, value="সর্বমোট").font = Font(bold=True)
ws.merge_cells("A1:N1"); ws.merge_cells("A3:N3"); ws.merge_cells("A4:N4")

# ===== AIS আমদানি =====
ws2 = wb.create_sheet("AIS Import Statement")
ws2["A1"] = "AIS Import Data — 01/01/2025 to 31/12/2025"
h2 = ["Sl","B/E No","B/E Date","Type","HS Code","Item Description","Qty","UOM",
      "Unit Price","CIF Value (USD)","Exchange Rate","Assessable Value BDT",
      "CD Paid","VAT Paid","RD Amount","SD Amount","AT Amount","AIT Amount",
      "Total Duty & Tax","Supplier","Country of Origin"]
for i,h in enumerate(h2,start=1):
    ws2.cell(row=3,column=i,value=h).font = Font(bold=True)

# ক্লাস্টার-১: 4 HS মিলে 120000+95000+88000+90000 = 393000 vs 370000 → 23000 অতিরিক্ত ✗
# ক্লাস্টার-২: 1500000+900000+600000 = 3000000 vs 3500000 → ঠিক আছে ✓
# 6006.32.00 : ক্লাস্টারে নেই → অননুমোদিত ✗
imports = [
 (1,"C-1001","12/01/2025","IM-4","5208.11.00","Grey Cotton Fabric",120000,"YDS",1.50,180000,121.5,21870000),
 (2,"C-1002","20/02/2025","IM-4","5209.11.00","Cotton Fabric Heavy",95000,"YDS",1.55,147250,121.8,17935050),
 (3,"C-1003","15/03/2025","IM-4","5407.61.00","PE Woven Fabric",88000,"YDS",1.20,105600,122.0,12883200),
 (4,"C-1004","10/05/2025","IM-4","9606.21.00","Plastic Button",1500000,"PCS",0.013,19500,122.5,2388750),
 (5,"C-1005","22/06/2025","IM-4","9607.11.00","Metal Zipper",900000,"PCS",0.045,40500,122.8,4973400),
 (6,"C-1006","18/07/2025","IM-4","5807.10.00","Woven Label",600000,"PCS",0.008,4800,123.0,590400),
 (7,"C-1007","05/09/2025","IM-4","6006.32.00","Knitted Fabric Dyed",30000,"KG",3.50,105000,123.5,12967500),
 (8,"C-1008","14/10/2025","IM-4","5513.11.00","Polyester Staple Fabric",90000,"YDS",1.35,121500,123.8,15041700),
 (9,"C-1009","20/11/2025","IM-4","5401.10.00","Sewing Thread Polyester",9500,"KG",3.15,29925,124.0,3710700),
]
RATES = {"5208.11.00":(5,0,0,15,5,5),"5209.11.00":(5,0,0,15,5,5),
         "5407.61.00":(5,0,0,15,5,5),"5513.11.00":(5,0,0,15,5,5),
         "9606.21.00":(10,3,0,15,5,5),"9607.11.00":(10,3,0,15,5,5),
         "5807.10.00":(10,3,0,15,5,5),"6006.32.00":(5,0,0,15,5,5),
         "5401.10.00":(5,0,0,15,5,5)}
r = 4
for row in imports:
    for i,v in enumerate(row,start=1):
        ws2.cell(row=r,column=i,value=v)
    hs, av = row[4], row[11]
    cd_r,rd_r,sd_r,vat_r,at_r,ait_r = RATES[hs]
    cd = av*cd_r/100; rd = av*rd_r/100
    sd = (av+cd+rd)*sd_r/100
    vat = (av+cd+rd+sd)*vat_r/100
    at = av*at_r/100; ait = av*ait_r/100
    for col, v in [(13,cd),(14,vat),(15,rd),(16,sd),(17,at),(18,ait),
                   (19,cd+rd+sd+vat+at+ait)]:
        ws2.cell(row=r,column=col,value=round(v))
    ws2.cell(row=r,column=20,value="Supplier Co")
    ws2.cell(row=r,column=21,value="China")
    r += 1

wb.save(OUT)
print(f"✅ তৈরি: {OUT}")
print("   ক্লাস্টার-১ (Merged): ৪টি কাপড়, প্রাপ্যতা ৩,৭০,০০০ YDS")
print("   ক্লাস্টার-২ (খালি সারি): ৩টি আনুষঙ্গিক, প্রাপ্যতা ৩৫,০০,০০০ PCS")
