"""
নিরীক্ষা প্রতিবেদন (.docx) পরীক্ষা
=====================================
চালান:  PYTHONPATH=backend python3 tests/test_audit_report_docx.py
"""
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))

from docx import Document                                   # noqa: E402
from docx.oxml.ns import qn                                 # noqa: E402

from knowledge.report_style import AuditProfile             # noqa: E402
from services.audit_report_docx import (                    # noqa: E402
    ReportContext, build_audit_report,
)

FAILS: list[str] = []
TMP = Path(tempfile.mkdtemp())


def check(label: str, ok: bool) -> None:
    print(f"  {'✅' if ok else '❌'} {label}")
    if not ok:
        FAILS.append(label)


def all_runs(doc):
    for p in doc.paragraphs:
        yield from p.runs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield from p.runs


def font_of(run) -> str:
    if run._element.rPr is None:
        return ""
    rf = run._element.rPr.find(qn("w:rFonts"))
    return rf.get(qn("w:cs")) if rf is not None else ""


def is_bengali(s: str) -> bool:
    return any(0x0980 <= ord(c) <= 0x09FF for c in s)


def make_ctx(**kw) -> ReportContext:
    base = dict(
        company="নমুনা টেক্সটাইল মিলস লিমিটেড",
        address="গাজীপুর",
        entity_label="সরাসরি রপ্তানিমুখী",
        bond_license="ঢাকা/বন্ড/১২৩",
        period_from="০১/০৭/২০২৩", period_to="৩০/০৬/২০২৪",
        profile=AuditProfile(),
        summary={
            "দাবি ২ — প্রাপ্যতার অতিরিক্ত আমদানি (BDT)": 1542936.0,
            "সর্বমোট রাজস্ব দাবি (BDT)": 1542936.0,
        },
        findings=[{"kind": "excess_import_demand", "para": "৩৬",
                   "quantity": "১৪১.২৯১", "unit": "মেঃ টন",
                   "amount": 1542936.0}],
    )
    base.update(kw)
    return ReportContext(**base)


print("== ১. দুই ফন্টেই নথি তৈরি হয় ==")
paths = {}
for f in ("nikosh", "sutonnymj"):
    p = build_audit_report(make_ctx(), TMP / f"r_{f}.docx", font=f)
    paths[f] = p
    check(f"{f} নথি তৈরি", p.exists() and p.stat().st_size > 10000)

print("== ২. Nikosh — লেখা ইউনিকোডেই থাকে ==")
d = Document(str(paths["nikosh"]))
runs = list(all_runs(d))
check("সব রানে Nikosh", all(font_of(r) in ("Nikosh", "")
                            for r in runs))
check("বাংলা লেখা অবিকৃত আছে",
      any("নিরীক্ষা প্রতিবেদন" in r.text for r in runs))

print("== ৩. ★ SutonnyMJ — ইউনিকোড বাংলা কোথাও ফাঁস হয় নাই ==")
d = Document(str(paths["sutonnymj"]))
runs = list(all_runs(d))
leak = [r.text for r in runs
        if font_of(r) == "SutonnyMJ" and is_bengali(r.text)]
check("বিজয় রানে ইউনিকোড বাংলা নাই", not leak)
eng = [r.text for r in runs
       if font_of(r) == "Times New Roman" and is_bengali(r.text)]
check("ইংরেজি রানে বাংলা অক্ষর নাই", not eng)
check("দুই ফন্টই ব্যবহৃত",
      {"SutonnyMJ", "Times New Roman"} <= {font_of(r) for r in runs})

print("== ৪. complex-script ক্ষেত্র বসানো আছে ==")
ok = True
for r in runs:
    if r._element.rPr is None or not r.text:
        continue
    if r._element.rPr.find(qn("w:szCs")) is None:
        ok = False
        break
check("প্রতিটি রানে w:szCs আছে", ok)

print("== ৫. আপত্তির অঙ্ক নথিতে আসে ==")
d = Document(str(paths["nikosh"]))
txt = "\n".join(r.text for r in all_runs(d))
check("দাবির অঙ্ক আছে", "১৫,৪২,৯৩৬" in txt or "15,42,936" in txt)
check("মতামত অংশ আছে", "প্রস্তাব ও মতামত" in txt)

print("== ৬. ★ আপত্তি না থাকিলে দাবি বানায় না ==")
clean = make_ctx(summary={}, findings=[])
p = build_audit_report(clean, TMP / "clean.docx", font="nikosh")
txt = "\n".join(r.text for r in all_runs(Document(str(p))))
check("‘আপত্তি উদ্ঘাটিত হয় নাই’ লেখা আছে",
      "আপত্তি উদ্ঘাটিত হয় নাই" in txt)
check("কোনো দাবির ছক নাই", "সর্বমোট" not in txt)

print("== ৭. ★ দলিল না থাকিলে ‘নির্ণয় করা যায় নাই’ লেখে ==")
miss = make_ctx(missing_docs=["bond_register"])
p = build_audit_report(miss, TMP / "miss.docx", font="nikosh")
txt = "\n".join(r.text for r in all_runs(Document(str(p))))
check("ঘাটতির অনুচ্ছেদ আছে", "দলিলের ঘাটতি" in txt)
check("‘নির্ণয় করা সম্ভব হয় নাই’ আছে", "নির্ণয় করা সম্ভব হয় নাই" in txt)
check("‘দাবি উত্থাপন করা হয় নাই’ আছে", "দাবি উত্থাপন করা হয় নাই" in txt)

print("== ৮. প্রাসঙ্গিকতা — ইপিজেডে ইউপি-অনুচ্ছেদ বাদ ==")
epz = make_ctx(profile=AuditProfile(is_epz=True), findings=[])
p = build_audit_report(epz, TMP / "epz.docx", font="nikosh")
txt = "\n".join(r.text for r in all_runs(Document(str(p))))
check("ইউপি-অনুচ্ছেদ নাই", "ইউটিলাইজেশন পারমিশন" not in txt)

print()
if FAILS:
    print(f"RESULT: ❌ {len(FAILS)}টি ব্যর্থ")
    for f in FAILS:
        print("   -", f)
    sys.exit(1)
print("RESULT: ALL PASS ✅")
