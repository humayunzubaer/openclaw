"""
বাস্তবসম্মত প্রাপ্যতা শীট + AIS আমদানি শীট তৈরি (ব্যবহারকারীর বর্ণনা অনুযায়ী)
"""
import openpyxl
from openpyxl.styles import Font, Alignment
from pathlib import Path

OUT = Path(__file__).parent / "sample_entitlement_real.xlsx"
wb = openpyxl.Workbook()

# ================= Sheet 1: প্রাপ্যতা শীট =================
ws = wb.active
ws.title = "প্রাপ্যতা"

ws["A1"] = "গণপ্রজাতন্ত্রী বাংলাদেশ সরকার"
ws["A2"] = "কাস্টমস বন্ড কমিশনারেট, ঢাকা"
ws["A3"] = "প্রতিষ্ঠান: ABC Apparels Ltd. | বন্ড লাইসেন্স নং: DHK/BOND/1234"
ws["A4"] = "কাঁচামালের প্রাপ্যতা নির্ধারণ সংক্রান্ত"
ws["A5"] = ""

# ★ মূল কলাম — শিরোনামে মেয়াদ উল্লেখ
PERIOD_HEADER = "০১.০১.২০২৫ থেকে ৩১.১২.২০২৫ পর্যন্ত সময়ের জন্য প্রদত্ত/প্রস্তাবিত আমদানি প্রাপ্যতা"

headers = [
    "ক্রমিক নং",
    "এইচ.এস কোড",
    "কাঁচামালের বিবরণ",
    "বাণিজ্যিক নাম",
    "একক",
    PERIOD_HEADER,
    "একক মূল্য (USD)",
    "CD %", "RD %", "SD %", "VAT %", "AT %", "AIT %",
]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=6, column=i, value=h)
    c.font = Font(bold=True, size=9)
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[6].height = 60
ws.column_dimensions["F"].width = 30

rows = [
    [1, "5208.11.00", "Woven Fabrics of Cotton, Unbleached", "Grey Cotton Fabric", "YDS", 150000, 1.85, 5, 0, 0, 15, 5, 5],
    [2, "5407.61.00", "Woven Fabrics of Polyester Filament", "PE Woven Fabric", "YDS", 220000, 1.20, 5, 0, 0, 15, 5, 5],
    [3, "5402.33.00", "Textured Yarn of Polyester", "Poly Textured Yarn", "KG", 45000, 2.40, 5, 0, 0, 15, 5, 5],
    [4, "5806.32.00", "Narrow Woven Fabrics of Man-made Fibre", "Elastic Tape", "MTR", 80000, 0.15, 10, 3, 0, 15, 5, 5],
    [5, "9606.21.00", "Buttons of Plastics", "Plastic Button", "PCS", 2500000, 0.012, 10, 3, 0, 15, 5, 5],
    [6, "4819.10.00", "Cartons, Boxes of Corrugated Paper", "Export Carton", "PCS", 65000, 0.45, 15, 3, 0, 15, 5, 5],
    [7, "3919.10.00", "Self-adhesive Plates of Plastics", "Adhesive Tape", "PCS", 12000, 0.85, 10, 3, 0, 15, 5, 5],
    [8, "5401.10.00", "Sewing Thread of Synthetic Filaments", "Poly Sewing Thread", "KG", 18000, 3.10, 5, 0, 0, 15, 5, 5],
]
r = 7
for row in rows:
    for i, v in enumerate(row, start=1):
        ws.cell(row=r, column=i, value=v)
    r += 1
ws.cell(row=r, column=3, value="সর্বমোট").font = Font(bold=True)

ws.merge_cells("A1:M1")
ws.merge_cells("A2:M2")

# ================= Sheet 2: AIS আমদানি =================
ws2 = wb.create_sheet("AIS Import Statement")
ws2["A1"] = "AIS Import Data — Period: 01/01/2025 to 31/12/2025"

h2 = ["Sl", "B/E No", "B/E Date", "Type", "HS Code", "Item Description",
      "Qty", "UOM", "Unit Price", "CIF Value (USD)", "Exchange Rate",
      "Assessable Value BDT", "CD Paid", "VAT Paid", "Supplier", "Country of Origin"]
for i, h in enumerate(h2, start=1):
    ws2.cell(row=3, column=i, value=h).font = Font(bold=True)

# পরিস্থিতি:
#  - 5208.11.00 : প্রাপ্যতা 150000, আমদানি 45000+52000+68000 = 165000 → ১৫০০০ অতিরিক্ত ✗
#  - 5407.61.00 : প্রাপ্যতা 220000, আমদানি 68000+91000 = 159000 → ঠিক আছে ✓
#  - 6006.32.00 : প্রাপ্যতায় নেই → অননুমোদিত ✗
#  - 5806.32.00 : প্রাপ্যতা 80000, আমদানি 35000+50000 = 85000 → ৫০০০ অতিরিক্ত ✗
#  - একটি বিল মেয়াদের বাইরে (2024) → সতর্কতা
imports = [
    [1, "C-102345", "12/01/2025", "IM-4", "5208.11.00", "Grey Cotton Fabric", 45000, "YDS", 1.85, 83250, 121.50, 10114875, 505744, 1593092, "Zhejiang Textile Co", "China"],
    [2, "C-102876", "28/01/2025", "IM-4", "5407.61.00", "PE Woven Fabric", 68000, "YDS", 1.22, 82960, 121.50, 10079640, 503982, 1587543, "Jiangsu Fabrics", "China"],
    [3, "C-103455", "15/02/2025", "IM-4", "5402.33.00", "Polyester Textured Yarn", 15000, "KG", 2.45, 36750, 121.80, 4476150, 223808, 705006, "Suzhou Yarn Ltd", "China"],
    [4, "C-104120", "03/03/2025", "IM-4", "9606.21.00", "Plastic Button Assorted", 850000, "PCS", 0.013, 11050, 122.00, 1348100, 134810, 244007, "Ningbo Button Mfg", "China"],
    [5, "C-104899", "22/04/2025", "IM-4", "5208.11.00", "Cotton Grey Fabric", 52000, "YDS", 1.90, 98800, 122.20, 12073360, 603668, 1901554, "Zhejiang Textile Co", "China"],
    [6, "C-105677", "11/05/2025", "IM-4", "6006.32.00", "Knitted Fabric Dyed", 30000, "KG", 3.50, 105000, 122.50, 12862500, 643125, 2025844, "Guangdong Knit", "China"],
    [7, "C-106234", "05/06/2025", "IM-4", "4819.10.00", "Corrugated Export Carton", 42000, "PCS", 0.46, 19320, 122.80, 2372496, 355874, 409265, "Local Pack BD", "Bangladesh"],
    [8, "C-107001", "19/07/2025", "IM-4", "5401.10.00", "Sewing Thread Polyester", 9500, "KG", 3.15, 29925, 123.00, 3680775, 184039, 579728, "Coats Thread", "India"],
    [9, "C-108445", "14/08/2025", "IM-4", "5208.11.00", "Grey Cotton Fabric Unbleached", 68000, "YDS", 1.88, 127840, 123.20, 15749888, 787494, 2480607, "Zhejiang Textile Co", "China"],
    [10, "C-109233", "22/09/2025", "IM-7", "5806.32.00", "Elastic Tape Narrow", 35000, "MTR", 0.16, 5600, 123.50, 691600, 69160, 118864, "Foshan Elastic", "China"],
    [11, "C-109901", "30/10/2025", "IM-4", "5806.32.00", "Narrow Elastic Webbing", 50000, "MTR", 0.17, 8500, 123.80, 1052300, 105230, 180827, "Foshan Elastic", "China"],
    [12, "C-110002", "18/11/2025", "IM-4", "3919.10.00", "Adhesive Tape Roll", 5500, "PCS", 0.88, 4840, 124.00, 600160, 60016, 103128, "Shanghai Tape", "China"],
    [13, "C-110877", "09/12/2025", "IM-4", "5407.61.00", "Polyester Woven Fabric", 91000, "YDS", 1.25, 113750, 124.30, 14139125, 706956, 2226913, "Jiangsu Fabrics", "China"],
    # ★ মেয়াদের বাইরে — ২০২৪ সালের বিল
    [14, "C-098221", "15/11/2024", "IM-4", "5402.33.00", "Poly Textured Yarn DTY", 22000, "KG", 2.38, 52360, 119.50, 6257020, 312851, 985481, "Suzhou Yarn Ltd", "China"],
]
r = 4
for row in imports:
    for i, v in enumerate(row, start=1):
        ws2.cell(row=r, column=i, value=v)
    r += 1
ws2.cell(row=r, column=6, value="Grand Total").font = Font(bold=True)

wb.save(OUT)
print(f"✅ তৈরি হয়েছে: {OUT}")
print(f"   প্রাপ্যতা কলাম শিরোনাম: {PERIOD_HEADER}")
